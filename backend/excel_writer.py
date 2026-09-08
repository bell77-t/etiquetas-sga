"""
excel_writer.py - Editor Bidireccional de Productos en Base.xlsx / Base_Actualizada.xlsx
Permite modificar pictogramas, frases H/P, advertencia y unidad de medida desde la web y persistir en Excel.
"""

import shutil
import threading
from pathlib import Path
from typing import Dict, Any, List, Optional
import openpyxl

_EXCEL_WRITE_LOCK = threading.RLock()

def update_product_in_excel(base_path: Path, update_data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Actualiza un producto en Base_Actualizada.xlsx o Base.xlsx.
    update_data: {
        "codigo": "1ACABA01",
        "nombre": "ABAMECAL 1.8 EC",
        "palabra_advertencia": "PELIGRO",
        "frase_h": "...",
        "frase_p": "...",
        "um": "LITRO",
        "pictogramas": ["GHS06", "GHS09"]
    }
    """
    if not base_path.exists():
        raise FileNotFoundError(f"No se encontró el archivo base: {base_path}")

    with _EXCEL_WRITE_LOCK:
        return _update_product_in_excel_locked(base_path, update_data)


def _update_product_in_excel_locked(base_path: Path, update_data: Dict[str, Any]) -> Dict[str, Any]:
    # 1. Crear copia de seguridad antes de modificar
    backup_path = base_path.with_suffix(".xlsx.bak")
    shutil.copy2(base_path, backup_path)

    # 2. Cargar libro con openpyxl (preservando formato)
    wb = openpyxl.load_workbook(base_path, data_only=False)
    
    # Seleccionar hoja SGA
    target_sheet = None
    for s in wb.sheetnames:
        if "SGA" in s.upper():
            target_sheet = wb[s]
            break
    if target_sheet is None:
        target_sheet = wb.active

    ws = target_sheet

    # 3. Mapear encabezados
    header_row = 1
    col_map = {}
    for col in range(1, ws.max_column + 1):
        val = ws.cell(header_row, col).value
        if val:
            norm = str(val).strip().upper()
            col_map[norm] = col

    col_cod = col_map.get("CODIGO")
    col_nom = col_map.get("NOMBRE DEL PRODUCTO")
    col_um = col_map.get("U/M")
    col_fh = col_map.get("FRASES H")
    col_fp = col_map.get("FRASES P")
    col_adv = col_map.get("PALABRA DE ADVERTENCIA")
    col_pig1 = col_map.get("PIG1")
    col_pig2 = col_map.get("PIG2")
    col_pig3 = col_map.get("PIG3")
    col_pig4 = col_map.get("PIG4")

    # 4. Buscar la fila correspondiente
    target_row = None
    search_code = str(update_data.get("codigo", "")).strip().upper()
    search_name = str(update_data.get("nombre", "")).strip().upper()

    for row in range(2, ws.max_row + 1):
        row_cod = str(ws.cell(row, col_cod).value or "").strip().upper() if col_cod else ""
        row_nom = str(ws.cell(row, col_nom).value or "").strip().upper() if col_nom else ""

        if search_code and row_cod == search_code:
            target_row = row
            break
        elif search_name and (row_nom == search_name or search_name in row_nom):
            target_row = row
            break

    if not target_row:
        raise ValueError(f"No se encontró el producto '{search_name or search_code}' en la hoja '{ws.title}'.")

    # 5. Aplicar cambios
    if "palabra_advertencia" in update_data and col_adv:
        ws.cell(target_row, col_adv).value = str(update_data["palabra_advertencia"]).strip().upper()

    if "frase_h" in update_data and col_fh:
        ws.cell(target_row, col_fh).value = str(update_data["frase_h"]).strip()

    if "frase_p" in update_data and col_fp:
        ws.cell(target_row, col_fp).value = str(update_data["frase_p"]).strip()

    if "um" in update_data and col_um and update_data["um"]:
        ws.cell(target_row, col_um).value = str(update_data["um"]).strip().upper()

    # Asignar pictogramas (PIG1, PIG2, PIG3, PIG4)
    if "pictogramas" in update_data:
        p_list = update_data["pictogramas"] or []
        pig_cols = [col_pig1, col_pig2, col_pig3, col_pig4]
        for i, c_idx in enumerate(pig_cols):
            if c_idx:
                if i < len(p_list) and p_list[i]:
                    val = str(p_list[i]).strip()
                    # Si solo viene GHS05, guardar como 'GHS05'
                    ws.cell(target_row, c_idx).value = val
                else:
                    ws.cell(target_row, c_idx).value = "SIN FOTO"

    # 6. Guardar cambios en el archivo Excel
    wb.save(base_path)
    wb.close()

    return {
        "success": True,
        "row": target_row,
        "sheet": ws.title,
        "message": f"Producto '{search_name or search_code}' actualizado exitosamente en Excel."
    }
