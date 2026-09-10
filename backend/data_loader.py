import os
import io
import re
import datetime
from pathlib import Path
from typing import Dict, List, Any, Optional
import openpyxl

from .ghs_engine import infer_ghs_pictograms
from .validation import validate_product

PROJECT_ROOT = Path(__file__).resolve().parent.parent
BUNDLED_CATALOG_DIR = PROJECT_ROOT / "data" / "catalog"
BUNDLED_PICTOGRAM_DIR = PROJECT_ROOT / "static" / "picto"
LEGACY_DATA_DIR = Path.home() / "Downloads" / "Alejandra" / "Alejandra"

DEFAULT_SEARCH_PATHS = [
    Path(os.environ["SGA_DATA_DIR"]) if os.environ.get("SGA_DATA_DIR") else PROJECT_ROOT,
    PROJECT_ROOT,
    Path.cwd(),
    # Compatibilidad temporal con la instalación original. Use SGA_DATA_DIR en instalaciones nuevas.
    Path(r"C:\Users\alexc\Downloads\Alejandra\Alejandra"),
]

def find_data_directory(custom_path: Optional[str] = None) -> Path:
    """Busca el directorio que contiene ``aplicacion.xlsm`` (o usa SGA_DATA_DIR)."""
    if custom_path and os.path.exists(custom_path):
        p = Path(custom_path)
        if (p / "aplicacion.xlsm").exists():
            return p

    candidates = []
    if custom_path:
        candidates.append(Path(custom_path).expanduser())
    if os.environ.get("SGA_DATA_DIR"):
        candidates.append(Path(os.environ["SGA_DATA_DIR"]).expanduser())

    # El Excel puede vivir junto al proyecto, en data/ o en el catalogo incluido.
    candidates.extend([
        PROJECT_ROOT,
        PROJECT_ROOT / "data",
        BUNDLED_CATALOG_DIR,
        Path.cwd(),
        Path.cwd() / "data",
        Path.cwd() / "data" / "catalog",
        LEGACY_DATA_DIR,
    ])

    seen = set()
    for path in candidates:
        try:
            path = path.resolve()
        except OSError:
            continue
        if path in seen or not path.is_dir():
            continue
        seen.add(path)
        if (path / "aplicacion.xlsm").is_file():
            return path

    return PROJECT_ROOT


def read_file_bytes_non_blocking(filepath: Path) -> io.BytesIO:
    """Lee el archivo en un búfer en memoria permitiendo lectura compartida si está abierto en Excel."""
    if not filepath.exists():
        raise FileNotFoundError(f"No se encontró el archivo: {filepath}")
    
    with open(filepath, "rb") as f:
        data = f.read()
    return io.BytesIO(data)


def normalize_text(text: Any) -> str:
    """Normaliza texto para búsquedas y cruces insensibles a mayúsculas/espacios."""
    if text is None:
        return ""
    text_str = str(text).strip().upper()
    text_str = re.sub(r"\s+", " ", text_str)
    return text_str


def get_pictogram_catalog(picto_dir: Path) -> Dict[str, str]:
    """Escanea la carpeta Picto y crea un mapa de búsqueda por código GHS y nombre."""
    catalog = {}
    if not picto_dir.exists() or not picto_dir.is_dir():
        return catalog

    for file in picto_dir.glob("*.jpg"):
        fname = file.name
        norm = normalize_text(fname)
        catalog[norm] = fname
        match = re.search(r"(GHS\d{2})", norm)
        if match:
            ghs_code = match.group(1)
            catalog[ghs_code] = fname
            catalog[f"{ghs_code}.JPG"] = fname

    return catalog


def resolve_pictogram(pic_val: Any, picto_catalog: Dict[str, str], picto_dir: Path) -> Dict[str, Any]:
    """Resuelve la URL y metadatos de un pictograma para la etiqueta SGA."""
    if not pic_val:
        return {"has_image": False, "url": None, "label": "SIN IMAGEN", "code": None}

    val_str = str(pic_val).strip()
    norm = normalize_text(val_str)

    # Detectar si contiene 'SIN FOTO', 'SIN F', 'NO TIENE', 'N/A', etc.
    if any(k in norm for k in ["SIN FOTO", "SIN F", "NO TIENE", "NO APLICA", "N/A", "NINGUNO"]):
        return {"has_image": False, "url": None, "label": "SIN IMAGEN", "code": None}

    # Extraer código GHS (ej: GHS01, GHS06, etc.)
    match = re.search(r"(GHS\d{2})", norm)
    ghs_code = match.group(1) if match else None

    # Buscar archivo físico
    target_filename = None
    if ghs_code and (picto_dir / f"{ghs_code}.jpg").exists():
        target_filename = f"{ghs_code}.jpg"
    elif norm in picto_catalog and (picto_dir / picto_catalog[norm]).exists():
        target_filename = picto_catalog[norm]
    elif ghs_code and ghs_code in picto_catalog and (picto_dir / picto_catalog[ghs_code]).exists():
        target_filename = picto_catalog[ghs_code]
    else:
        # Búsqueda parcial
        for cat_norm, fname in picto_catalog.items():
            if (cat_norm in norm or norm in cat_norm) and (picto_dir / fname).exists():
                target_filename = fname
                break

    if target_filename:
        return {
            "has_image": True,
            "url": f"/picto/{target_filename}",
            "filename": target_filename,
            "label": val_str,
            "code": ghs_code or val_str
        }

    return {"has_image": False, "url": None, "label": "SIN IMAGEN", "code": None}


def limit_ghs_phrases(text: Any, max_items: int = 3, prefix_char: str = 'H') -> str:
    """Limita frases H o P a un máximo de `max_items` (por defecto 3) para legibilidad y diseño."""
    if not text or not str(text).strip():
        return ""
    raw = str(text).strip()
    
    # 1. Si viene con saltos de línea explícitos
    lines = [line.strip() for line in raw.split('\n') if line.strip()]
    if len(lines) > 1:
        return '\n'.join(lines[:max_items])
    
    # 2. Si viene separado por múltiples espacios (ej: 'H319 ...       H335 ...')
    clean_text = re.sub(r'[ \t]{2,}', '\n', raw)
    lines2 = [l.strip() for l in clean_text.split('\n') if l.strip()]
    if len(lines2) > 1:
        return '\n'.join(lines2[:max_items])
    
    # 3. Si sigue siendo una sola línea con múltiples códigos Hxxx o Pxxx (evitando partir frases compuestas Pxxx + Pxxx)
    parts = re.split(rf'(?<!\+)\s*(?=\b{prefix_char}\d{{3}})', raw)
    valid_parts = [p.strip() for p in parts if p.strip()]
    if len(valid_parts) > 1:
        return '\n'.join(valid_parts[:max_items])
        
    return raw


def load_base_catalog(base_path: Path, picto_dir: Path) -> Dict[str, Dict[str, Any]]:
    """Carga el maestro de productos desde Base.xlsx (hoja SGA)."""
    buf = read_file_bytes_non_blocking(base_path)
    wb = openpyxl.load_workbook(buf, data_only=True, read_only=True)
    
    # Buscar hoja que contenga 'SGA' (ej. 'SGA AJ', 'SGA', etc.)
    sheet_name = None
    for s in wb.sheetnames:
        if "SGA" in s.upper():
            sheet_name = s
            break
    if not sheet_name:
        sheet_name = wb.sheetnames[0]
    ws = wb[sheet_name]

    picto_catalog = get_pictogram_catalog(picto_dir)
    products = {}

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return products

    # Buscar fila de encabezados
    header_idx = 0
    header_map = {}
    for idx, row in enumerate(rows[:10]):
        row_str = [normalize_text(c) for c in row if c is not None]
        if "CODIGO" in row_str or "NOMBRE DEL PRODUCTO" in row_str:
            header_idx = idx
            for col_idx, col_name in enumerate(row):
                if col_name is not None:
                    header_map[normalize_text(col_name)] = col_idx
            break

    def get_val(row_tuple, col_name, default=""):
        idx = header_map.get(normalize_text(col_name))
        if idx is not None and idx < len(row_tuple) and row_tuple[idx] is not None:
            return row_tuple[idx]
        return default

    for row in rows[header_idx + 1:]:
        if not any(row):
            continue
        
        prod_name = get_val(row, "Nombre del Producto")
        if not prod_name or str(prod_name).strip() == "":
            continue

        codigo = str(get_val(row, "Codigo", "")).strip()
        um = str(get_val(row, "U/m", "KILO")).strip().upper()
        if not um:
            um = "KILO"
            
        raw_frase_h = str(get_val(row, "FRASES H", "")).strip()
        raw_frase_p = str(get_val(row, "FRASES P", "")).strip()
        
        # Limitar a máximo 3 frases H y 3 frases P
        frase_h = limit_ghs_phrases(raw_frase_h, max_items=3, prefix_char='H')
        frase_p = limit_ghs_phrases(raw_frase_p, max_items=3, prefix_char='P')
        
        palabra_adv = str(get_val(row, "PALABRA DE ADVERTENCIA", "")).strip().upper()
        
        pig1_val = get_val(row, "PIG1", "SIN FOTO")
        pig2_val = get_val(row, "PIG2", "SIN FOTO")
        pig3_val = get_val(row, "PIG3", "SIN FOTO")
        pig4_val = get_val(row, "PIG4", "SIN FOTO")


        pictograms = [
            resolve_pictogram(pig1_val, picto_catalog, picto_dir),
            resolve_pictogram(pig2_val, picto_catalog, picto_dir),
            resolve_pictogram(pig3_val, picto_catalog, picto_dir),
            resolve_pictogram(pig4_val, picto_catalog, picto_dir),
        ]

        has_real_picto = any(p.get("has_image") for p in pictograms)
        inferred_list = []
        if not has_real_picto:
            inferred = infer_ghs_pictograms(frase_h, palabra_adv)
            if inferred:
                inferred_list = inferred
                new_pictos = []
                for p in inferred[:4]:
                    new_pictos.append({
                        "has_image": True,
                        "url": f"/picto/{p['filename']}",
                        "filename": p["filename"],
                        "label": p["name"],
                        "code": p["code"],
                        "inferred": True,
                        "reason": p.get("reason", "")
                    })
                while len(new_pictos) < 4:
                    new_pictos.append({"has_image": False, "url": None, "label": "SIN IMAGEN", "code": None})
                pictograms = new_pictos

        prod_info = {
            "codigo": codigo,
            "nombre": str(prod_name).strip(),
            "um": um,
            "frase_h": frase_h,
            "frase_p": frase_p,
            "palabra_advertencia": palabra_adv,
            "pictogramas": pictograms,
            "has_inferred_pictos": bool(inferred_list),
            "inferred_reasons": [p.get("reason") for p in inferred_list],
            "tiene_sga": str(get_val(row, "Tiene SGA", "")).strip(),
            "tiene_hds": str(get_val(row, "TIENE HOJA DE SEGURIDAD", "")).strip(),
        }
        prod_info["faltantes_sga"] = validate_product(prod_info)
        prod_info["estado_ficha"] = "COMPLETA" if not prod_info["faltantes_sga"] else "INCOMPLETA"

        norm_key = normalize_text(prod_name)
        products[norm_key] = prod_info
        if codigo:
            products[f"COD:{codigo.upper()}"] = prod_info

    return products


def format_date_str(val: Any) -> Dict[str, str]:
    """Formatea fecha a formato ISO y formato latino legible."""
    if isinstance(val, (datetime.datetime, datetime.date)):
        iso_val = val.strftime("%Y-%m-%d")
        display_val = val.strftime("%d/%m/%Y")
        return {"iso": iso_val, "display": display_val}
    elif isinstance(val, str):
        val_str = val.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y"):
            try:
                dt = datetime.datetime.strptime(val_str, fmt)
                return {"iso": dt.strftime("%Y-%m-%d"), "display": dt.strftime("%d/%m/%Y")}
            except ValueError:
                pass
        return {"iso": val_str, "display": val_str}
    return {"iso": "", "display": "N/A"}


def load_applications(
    app_path: Path, 
    base_catalog: Dict[str, Dict[str, Any]],
    sheet_name: str = "Data"
) -> List[Dict[str, Any]]:
    """Carga y procesa la programación semanal desde aplicacion.xlsm para la hoja especificada."""
    buf = read_file_bytes_non_blocking(app_path)
    wb = openpyxl.load_workbook(buf, data_only=True, read_only=True)
    
    target_sheet = sheet_name if sheet_name in wb.sheetnames else ("Data" if "Data" in wb.sheetnames else wb.sheetnames[0])
    ws = wb[target_sheet]

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header_idx = -1
    header_map = {}
    for idx, row in enumerate(rows[:15]):
        row_str = [normalize_text(c) for c in row if c is not None]
        if "PRODUCTO" in row_str or "FECHA" in row_str or "LITROS AGUA" in row_str:
            header_idx = idx
            for col_idx, col_name in enumerate(row):
                if col_name is not None:
                    cleaned = normalize_text(col_name)
                    header_map[cleaned] = col_idx
            break

    if header_idx == -1:
        header_idx = 0
        for col_idx, col_name in enumerate(rows[0]):
            if col_name:
                header_map[normalize_text(col_name)] = col_idx

    def get_val(row_tuple, *col_aliases, default=None):
        for alias in col_aliases:
            idx = header_map.get(normalize_text(alias))
            if idx is not None and idx < len(row_tuple) and row_tuple[idx] is not None:
                return row_tuple[idx]
            norm_alias = normalize_text(alias)
            for h_key, h_idx in header_map.items():
                if norm_alias in h_key and h_idx < len(row_tuple) and row_tuple[h_idx] is not None:
                    return row_tuple[h_idx]
        return default

    applications = []
    item_counter = 0

    # Variables de contexto para herencia entre filas (típico en hojas ALZ y R-S-L)
    curr_date_raw = None
    default_cultivo = "Alstroemeria" if target_sheet == "ALZ" else ""
    curr_cultivo = default_cultivo
    curr_bloque = ""
    curr_camas = ""
    curr_litros = 0.0

    for row_num, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(row):
            continue

        # Saltar filas de totales (ej: 'Total 26/02/2024')
        if str(row[0] or "").strip().upper().startswith("TOTAL"):
            continue

        raw_prod = get_val(row, "PRODUCTO", "Producto", default="")
        
        # Actualizar contexto de fecha si viene en la fila
        d_cand = get_val(row, "Fecha", "FECHA")
        if d_cand is not None and str(d_cand).strip():
            curr_date_raw = d_cand

        # Actualizar contexto de cultivo
        c_cand = get_val(row, "Cultivo")
        if c_cand is not None and str(c_cand).strip() and not str(c_cand).strip().upper().startswith("TOTAL"):
            curr_cultivo = str(c_cand).strip()

        # Actualizar contexto de bloque
        b_cand = get_val(row, "Bloque")
        if b_cand is not None and str(b_cand).strip():
            curr_bloque = str(b_cand).strip()

        # Actualizar contexto de camas
        camas_cand = get_val(row, "Camas")
        if camas_cand is not None and str(camas_cand).strip():
            curr_camas = camas_cand

        # Actualizar contexto de litros
        l_cand = get_val(row, "Litros Agua", "Litros")
        if l_cand is not None:
            try:
                l_num = float(l_cand)
                if l_num > 0:
                    curr_litros = l_num
            except (ValueError, TypeError):
                pass

        # Si no hay producto válido en esta fila, continuar (era solo fila de encabezado de bloque)
        if not raw_prod or str(raw_prod).strip() in ("", "(en blanco)", "0", "None") or normalize_text(raw_prod) == "PRODUCTO":
            continue

        prod_name = str(raw_prod).strip()
        date_info = format_date_str(curr_date_raw)
        cultivo = curr_cultivo or default_cultivo
        bloque = curr_bloque
        camas = curr_camas

        reentrada = get_val(row, "Reentrada", default=0)
        categoria = get_val(row, "Categoria", "Categoría", default="")
        observaciones = str(get_val(row, "Observaciones", default="") or "").strip()
        semana = get_val(row, "Semana", default="")

        litros_total = curr_litros
        # Normalizar si el volumen está expresado en cc/ml (> 5,000 cc para cama/bloque estándar)
        if litros_total > 5000:
            litros_total = round(litros_total / 1000.0, 2)

        dosis_raw = get_val(row, "Dosis gr ó cc x Litro", "Dosis gr  cc x Litro", "Dosis", default=0)
        try:
            dosis = float(dosis_raw) if dosis_raw is not None else 0.0
        except (ValueError, TypeError):
            dosis = 0.0

        total_raw = get_val(row, "Total gramo ó cc", "Total gramo  cc", "Total", default=None)
        try:
            total_producto = float(total_raw) if total_raw is not None else (litros_total * dosis)
        except (ValueError, TypeError):
            total_producto = litros_total * dosis

        norm_pname = normalize_text(prod_name)
        base_info = base_catalog.get(norm_pname)
        
        if not base_info:
            for b_key, b_val in base_catalog.items():
                if b_key in norm_pname or norm_pname in b_key:
                    base_info = b_val
                    break

        if not base_info:
            base_info = {
                "codigo": "",
                "nombre": prod_name,
                "um": "KILO" if "SULFATO" in norm_pname or "NITRATO" in norm_pname or "ACIDO" not in norm_pname else "LITRO",
                "frase_h": "Información SGA no registrada en Base.xlsx para este producto.",
                "frase_p": "Consulte la Ficha de Datos de Seguridad y use equipo de protección personal adecuado.",
                "palabra_advertencia": "ATENCIÓN",
                "pictogramas": [
                    {"has_image": False, "url": None, "label": "SIN IMAGEN", "code": None},
                    {"has_image": False, "url": None, "label": "SIN IMAGEN", "code": None},
                    {"has_image": False, "url": None, "label": "SIN IMAGEN", "code": None},
                    {"has_image": False, "url": None, "label": "SIN IMAGEN", "code": None},
                ],
                "tiene_sga": "NO",
                "tiene_hds": "NO",
            }

        sector_bloque = ""
        if cultivo and bloque:
            sector_bloque = f"{cultivo} - Bloque {bloque}"
        elif cultivo:
            sector_bloque = cultivo
        elif bloque:
            sector_bloque = f"Bloque {bloque}"
        else:
            sector_bloque = "General"

        item_counter += 1
        app_item = {
            "id": f"app-{item_counter}",
            "row_num": row_num,
            "fecha": date_info,
            "semana": semana,
            "cultivo": cultivo,
            "bloque": bloque,
            "sector_bloque": sector_bloque,
            "camas": camas,
            "reentrada": reentrada,
            "categoria": categoria,
            "producto": prod_name,
            "litros_total": litros_total,
            "dosis": dosis,
            "total_producto": total_producto,
            "observaciones": observaciones,
            "base_info": base_info,
        }

        app_item["etiquetas"] = generate_tank_labels(app_item)
        applications.append(app_item)

    return applications


def generate_tank_labels(app_item: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Genera las etiquetas desglosadas por tanques de 1,000 Litros y colitas."""
    volumen_total = app_item["litros_total"]
    dosis = app_item["dosis"]
    base_info = app_item["base_info"]
    um = base_info["um"]

    labels = []

    if volumen_total <= 0:
        return labels

    if volumen_total < 1000:
        cant_dosificar = volumen_total * dosis
        label = {
            "id": f"{app_item['id']}-T1",
            "tanque_num": 1,
            "total_tanques": 1,
            "tipo_tanque": "Tanque Único",
            "es_colita": False,
            "litros_tanque": volumen_total,
            "dosis": dosis,
            "cantidad_dosificar": cant_dosificar,
            "unidad": um,
            "producto": app_item["producto"],
            "sector_bloque": app_item["sector_bloque"],
            "reentrada": app_item["reentrada"],
            "categoria": app_item["categoria"],
            "fecha": app_item["fecha"],
            "base_info": base_info,
            "litros_total_lote": volumen_total,
        }
        labels.append(label)
        return labels

    n_completos = int(volumen_total // 1000)
    residuo = volumen_total % 1000
    total_etiquetas = n_completos + (1 if residuo > 0 else 0)

    for i in range(1, n_completos + 1):
        cant_dosificar = 1000.0 * dosis
        label = {
            "id": f"{app_item['id']}-T{i}",
            "tanque_num": i,
            "total_tanques": total_etiquetas,
            "tipo_tanque": f"Tanque {i} de {total_etiquetas} (1,000 L)",
            "es_colita": False,
            "litros_tanque": 1000.0,
            "dosis": dosis,
            "cantidad_dosificar": cant_dosificar,
            "unidad": um,
            "producto": app_item["producto"],
            "sector_bloque": app_item["sector_bloque"],
            "reentrada": app_item["reentrada"],
            "categoria": app_item["categoria"],
            "fecha": app_item["fecha"],
            "base_info": base_info,
            "litros_total_lote": volumen_total,
        }
        labels.append(label)

    if residuo > 0:
        cant_dosificar = residuo * dosis
        label = {
            "id": f"{app_item['id']}-T{total_etiquetas}",
            "tanque_num": total_etiquetas,
            "total_tanques": total_etiquetas,
            "tipo_tanque": f"Colita {total_etiquetas} de {total_etiquetas} ({residuo:.1f} L)".replace(".0 L", " L"),
            "es_colita": True,
            "litros_tanque": residuo,
            "dosis": dosis,
            "cantidad_dosificar": cant_dosificar,
            "unidad": um,
            "producto": app_item["producto"],
            "sector_bloque": app_item["sector_bloque"],
            "reentrada": app_item["reentrada"],
            "categoria": app_item["categoria"],
            "fecha": app_item["fecha"],
            "base_info": base_info,
            "litros_total_lote": volumen_total,
        }
        labels.append(label)

    return labels


class DataManager:
    """Administrador en memoria de los datos cargados con soporte de recarga en caliente."""
    def __init__(self, data_dir: Optional[str] = None):
        self.data_dir = find_data_directory(data_dir)
        self.base_catalog: Dict[str, Dict[str, Any]] = {}
        self.applications: List[Dict[str, Any]] = []
        self.last_loaded: Optional[datetime.datetime] = None
        self.available_dates: List[Dict[str, str]] = []
        self.available_products: List[str] = []
        self.load_all()

    def load_all(self):
        self.picto_dir = self.data_dir / "Picto"
        if not self.picto_dir.exists():
            self.picto_dir = BUNDLED_PICTOGRAM_DIR

        # Buscar archivo maestro de productos priorizando Base_Actualizada.xlsx
        base_candidates = [
            self.data_dir / "Base_Actualizada.xlsx",
            self.data_dir / "base_Actualizada.xlsx",
            self.data_dir / "Base_actualizada.xlsx",
            self.data_dir / "base_actualizada.xlsx",
            self.data_dir / "Base.xlsx",
            self.data_dir / "base.xlsx",
            BUNDLED_CATALOG_DIR / "Base_Actualizada.xlsx",
            BUNDLED_CATALOG_DIR / "Base.xlsx",
        ]
        base_path = None
        for cand in base_candidates:
            if cand.exists():
                base_path = cand
                break

        if not base_path:
            for f in self.data_dir.glob("*.xlsx"):
                if "base" in f.name.lower():
                    base_path = f
                    break

        app_path = self.data_dir / "aplicacion.xlsm"

        if not base_path or not base_path.exists():
            raise FileNotFoundError(f"No se encontró Base_Actualizada.xlsx ni Base.xlsx en {self.data_dir}")
        if not app_path.exists():
            raise FileNotFoundError(
                f"No se encontró aplicacion.xlsm en {self.data_dir}. "
                "Configure SGA_DATA_DIR con la carpeta que contiene los archivos Excel."
            )

        self.base_filename = base_path.name
        self.base_path = base_path
        self.base_catalog = load_base_catalog(base_path, self.picto_dir)

        # Extraer lista única de los 228 productos maestros desde Base_Actualizada.xlsx (Columna B: Nombre del Producto)
        seen_master_names = set()
        self.master_products = []
        for k, prod in self.base_catalog.items():
            if not k.startswith("COD:"):
                name = prod.get("nombre", "").strip()
                if name and name.upper() not in seen_master_names:
                    seen_master_names.add(name.upper())
                    self.master_products.append(prod)
        self.master_products.sort(key=lambda x: x["nombre"])

        # Cargar todos los programas disponibles en aplicacion.xlsm
        self.programs = {}
        sheet_meta = [
            ("Data", "Riego Sector 3 (Data)", "droplet"),
            ("ALZ", "Alstroemeria (ALZ)", "flower-2"),
            ("R-S-L", "Sector R-S-L", "layers"),
        ]

        buf = read_file_bytes_non_blocking(app_path)
        wb = openpyxl.load_workbook(buf, data_only=True, read_only=True)
        available_sheets = wb.sheetnames

        for s_id, s_name, s_icon in sheet_meta:
            if s_id in available_sheets:
                apps = load_applications(app_path, self.base_catalog, sheet_name=s_id)
                self.programs[s_id] = {
                    "id": s_id,
                    "name": s_name,
                    "icon": s_icon,
                    "count": len(apps),
                    "total_labels": sum(len(a["etiquetas"]) for a in apps),
                    "applications": apps,
                    "base_file": self.base_filename,
                    "is_mipe": False
                }

        # --- CARGAR BASE DE DATOS MIPE ---
        mipe_candidates = [
            self.data_dir / "BD_INFORMACION_ETIQUETAS_SGA_2025.xlsx",
            BUNDLED_CATALOG_DIR / "BD_INFORMACION_ETIQUETAS_SGA_2025.xlsx",
            self.data_dir / "BD INFORMACION ETIQUETAS SGA 2025.xlsx",
            BUNDLED_CATALOG_DIR / "BD INFORMACION ETIQUETAS SGA 2025.xlsx",
        ]
        for f_dir in [self.data_dir, BUNDLED_CATALOG_DIR]:
            if f_dir.exists():
                for candidate in f_dir.glob("*.xlsx"):
                    if "2025" in candidate.name or "mipe" in candidate.name.lower():
                        if candidate not in mipe_candidates:
                            mipe_candidates.append(candidate)

        mipe_path = None
        for cand in mipe_candidates:
            if cand.exists():
                mipe_path = cand
                break

        self.mipe_path = mipe_path
        self.mipe_catalog = {}
        self.mipe_master_products = []

        if mipe_path and mipe_path.exists():
            self.mipe_filename = mipe_path.name
            self.mipe_catalog = load_base_catalog(mipe_path, self.picto_dir)
            
            seen_mipe = set()
            for k, prod in self.mipe_catalog.items():
                if not k.startswith("COD:"):
                    name = prod.get("nombre", "").strip()
                    if name and name.upper() not in seen_mipe:
                        seen_mipe.add(name.upper())
                        self.mipe_master_products.append(prod)
            self.mipe_master_products.sort(key=lambda x: x["nombre"])

            # Crear aplicaciones interactivas para cada producto MIPE
            today_date = datetime.date.today().strftime("%Y-%m-%d")
            today_display = datetime.date.today().strftime("%d/%m/%Y")
            mipe_apps = []

            for i, p in enumerate(self.mipe_master_products, start=1):
                app_id = f"MIPE_{p.get('codigo') or i}"
                lbl = {
                    "id": f"{app_id}-T1",
                    "tanque_num": 1,
                    "total_tanques": 1,
                    "tipo_tanque": "Ficha MIPE SGA 2025 (Bodega / 1,000 L)",
                    "es_colita": False,
                    "litros_tanque": 1000.0,
                    "dosis": 1.0,
                    "cantidad_dosificar": 1000.0,
                    "unidad": p.get("um", "LITRO"),
                    "producto": p["nombre"],
                    "sector_bloque": "MIPE / BODEGA",
                    "reentrada": "0",
                    "categoria": "MIPE FITOSANITARIO",
                    "fecha": {"iso": today_date, "display": today_display},
                    "base_info": p,
                    "litros_total_lote": 1000.0,
                }
                app_item = {
                    "id": app_id,
                    "producto": p["nombre"],
                    "sector_bloque": "MIPE / BODEGA",
                    "fecha": {"iso": today_date, "display": today_display},
                    "dosis": 1.0,
                    "total_producto": 1000.0,
                    "litros_total": 1000.0,
                    "reentrada": "0",
                    "categoria": "MIPE FITOSANITARIO",
                    "observaciones": f"Ficha MIPE SGA 2025 - Código: {p.get('codigo', 'S/C')}",
                    "base_info": p,
                    "etiquetas": [lbl],
                    "is_mipe": True,
                }
                mipe_apps.append(app_item)

            self.programs["MIPE"] = {
                "id": "MIPE",
                "name": "MIPE (Catálogo Fitosanitarios)",
                "icon": "shield-check",
                "count": len(mipe_apps),
                "total_labels": len(mipe_apps),
                "applications": mipe_apps,
                "base_file": self.mipe_filename,
                "is_mipe": True,
            }

        # --- CARGAR PESTAÑA CATÁLOGO MIRFE COMPLETO (228 PRODUCTOS DE Base_Actualizada.xlsx) ---
        today_date = datetime.date.today().strftime("%Y-%m-%d")
        today_display = datetime.date.today().strftime("%d/%m/%Y")
        base_apps = []
        for i, p in enumerate(self.master_products, start=1):
            app_id = f"MIRFE_{p.get('codigo') or i}"
            lbl = {
                "id": f"{app_id}-T1",
                "tanque_num": 1,
                "total_tanques": 1,
                "tipo_tanque": "Ficha SGA MIRFE (Bodega / 1,000 L)",
                "es_colita": False,
                "litros_tanque": 1000.0,
                "dosis": 1.0,
                "cantidad_dosificar": 1000.0,
                "unidad": p.get("um", "KILO"),
                "producto": p["nombre"],
                "sector_bloque": "CATÁLOGO MIRFE / BODEGA",
                "reentrada": "0",
                "categoria": "FITOSANITARIO",
                "fecha": {"iso": today_date, "display": today_display},
                "base_info": p,
                "litros_total_lote": 1000.0,
            }
            app_item = {
                "id": app_id,
                "producto": p["nombre"],
                "sector_bloque": "CATÁLOGO MIRFE / BODEGA",
                "fecha": {"iso": today_date, "display": today_display},
                "dosis": 1.0,
                "total_producto": 1000.0,
                "litros_total": 1000.0,
                "reentrada": "0",
                "categoria": "FITOSANITARIO",
                "observaciones": f"Ficha SGA Catálogo MIRFE - Código: {p.get('codigo', 'S/C')}",
                "base_info": p,
                "etiquetas": [lbl],
                "is_base": True,
            }
            base_apps.append(app_item)

        self.programs["MIRFE"] = {
            "id": "MIRFE",
            "name": "MIRFE (Catálogo Fertirriego)",
            "icon": "droplet",
            "count": len(base_apps),
            "total_labels": len(base_apps),
            "applications": base_apps,
            "base_file": self.base_filename,
            "is_base": True,
        }

        if not hasattr(self, "current_program") or self.current_program not in self.programs:
            self.current_program = "MIPE" if "MIPE" in self.programs else ("MIRFE" if "MIRFE" in self.programs else list(self.programs.keys())[0])

        self.set_active_program(self.current_program)
        self.last_loaded = datetime.datetime.now()



    def set_active_program(self, program_id: str):
        if program_id not in self.programs:
            program_id = list(self.programs.keys())[0]
        self.current_program = program_id
        self.applications = self.programs[program_id]["applications"]

        dates_seen = set()
        dates_list = []
        products_seen = set()

        for item in self.applications:
            iso_d = item["fecha"]["iso"]
            if iso_d and iso_d not in dates_seen:
                dates_seen.add(iso_d)
                dates_list.append(item["fecha"])
            
            p_name = item["producto"]
            if p_name and p_name not in products_seen:
                products_seen.add(p_name)

        self.available_dates = sorted(dates_list, key=lambda x: x["iso"], reverse=True)
        self.available_products = sorted(list(products_seen))

    def get_program_data(self, program_id: Optional[str] = None) -> Dict[str, Any]:
        """Obtiene datos de un programa sin cambiar el estado compartido del servidor."""
        selected_id = program_id or self.current_program
        if selected_id not in self.programs:
            raise KeyError(selected_id)
        applications = self.programs[selected_id]["applications"]
        
        if selected_id in ["MIPE", "MIRFE", "BASE"]:
            available_dates = []
        else:
            dates = {item["fecha"]["iso"]: item["fecha"] for item in applications if item["fecha"]["iso"]}
            available_dates = sorted(dates.values(), key=lambda item: item["iso"], reverse=True)

            
        products = sorted({item["producto"] for item in applications if item["producto"]})
        return {
            "program_id": selected_id,
            "applications": applications,
            "available_dates": available_dates,
            "available_products": products,
        }


    def get_summary(self, program_id: Optional[str] = None) -> Dict[str, Any]:
        program_data = self.get_program_data(program_id)
        selected_id = program_data["program_id"]
        applications = program_data["applications"]
        total_apps = len(applications)
        total_labels = sum(len(a["etiquetas"]) for a in applications)
        # Exponer los dos catálogos normativos primarios: MIPE y MIRFE
        available_keys = [k for k in ["MIPE", "MIRFE"] if k in self.programs]
        if not available_keys:
            available_keys = list(self.programs.keys())

        programs_summary = [
            {
                "id": self.programs[k]["id"],
                "name": self.programs[k]["name"],
                "icon": self.programs[k]["icon"],
                "count": self.programs[k]["count"],
                "total_labels": self.programs[k]["total_labels"],
            }
            for k in available_keys
        ]
        
        is_mipe_selected = selected_id == "MIPE"
        master_products = self.mipe_master_products if is_mipe_selected else self.master_products
        base_file_name = getattr(self, "mipe_filename", "BD INFORMACIÓN ETIQUETAS SGA 2025.xlsx") if is_mipe_selected else getattr(self, "base_filename", "Base_Actualizada.xlsx")

        incomplete_list = []
        for p in master_products:
            missing_items = []
            if not any(pic.get("has_image") for pic in p.get("pictogramas", [])):
                missing_items.append("Sin pictogramas GHS")
            if not p.get("frase_h", "").strip():
                missing_items.append("Sin Frase H")
            if not p.get("frase_p", "").strip():
                missing_items.append("Sin Frase P")
            if not p.get("palabra_advertencia", "").strip():
                missing_items.append("Sin Palabra de Advertencia")
            
            if missing_items:
                incomplete_list.append({
                    "codigo": p.get("codigo") or "N/A",
                    "nombre": p["nombre"],
                    "um": p.get("um", "LITRO"),
                    "palabra_advertencia": p.get("palabra_advertencia", "PELIGRO"),
                    "faltantes": missing_items,
                    "has_picto": any(pic.get("has_image") for pic in p.get("pictogramas", [])),
                    "has_frase_h": bool(p.get("frase_h", "").strip()),
                    "has_frase_p": bool(p.get("frase_p", "").strip()),
                    "has_adv": bool(p.get("palabra_advertencia", "").strip()),
                })

        without_pictograms = sum(1 for item in incomplete_list if "Sin pictogramas GHS" in item["faltantes"])
        without_safety_text = sum(1 for item in incomplete_list if "Sin Frase H" in item["faltantes"] or "Sin Frase P" in item["faltantes"])

        return {
            "data_directory": str(self.data_dir),
            "base_file": base_file_name,
            "last_loaded": self.last_loaded.strftime("%Y-%m-%d %H:%M:%S") if self.last_loaded else None,
            "current_program": program_data["program_id"],
            "available_programs": programs_summary,
            "total_applications": total_apps,
            "total_labels": total_labels,
            "total_products_in_catalog": len(master_products),
            "master_products": [p["nombre"] for p in master_products],
            "master_products_data": master_products,
            "available_dates": program_data["available_dates"],
            "available_products": program_data["available_products"],
            "quality": {
                "without_pictograms": without_pictograms,
                "without_safety_text": without_safety_text,
                "total_incomplete": len(incomplete_list),
                "incomplete_products": incomplete_list,
            },
        }


