import openpyxl

wb_dl = openpyxl.load_workbook(r'C:\Users\alexc\Downloads\BD INFORMACIÓN ETIQUETAS SGA 2025 FF-AJ (1) (1).xlsx', data_only=True)
print("Downloads BD INFORMACION sheets:", wb_dl.sheetnames)
for s in wb_dl.sheetnames:
    ws = wb_dl[s]
    rows = [r for r in ws.iter_rows(values_only=True) if any(r)]
    print(f"  Hoja {s}: {len(rows)} filas no vacias")
