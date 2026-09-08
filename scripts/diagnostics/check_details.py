import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\alexc\Downloads\Alejandra\Alejandra\Base_Actualizada.xlsx', data_only=True)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    non_empty = [r for r in rows if any(r is not None and str(r).strip() != '' for r in r)]
    print(f"Hoja '{sheet}': Total filas={len(rows)}, Filas no vacias={len(non_empty)}")

# Revisar la columna Codigo en 'SGA AJ'
ws_sga = wb['SGA AJ']
rows_sga = list(ws_sga.iter_rows(values_only=True))
codigos = [r[0] for r in rows_sga[1:] if r[0] is not None and str(r[0]).strip() != '']
nombres = [r[1] for r in rows_sga[1:] if r[1] is not None and str(r[1]).strip() != '']
print(f"En 'SGA AJ': Codigos no vacios={len(codigos)}, Nombres no vacios={len(nombres)}")

# Revisar si hay duplicados de codigos
from collections import Counter
cod_counts = Counter([str(c).strip().upper() for c in codigos])
cod_dups = {k: v for k, v in cod_counts.items() if v > 1}
print(f"Codigos duplicados en 'SGA AJ': {len(cod_dups)}")
for k, v in cod_dups.items():
    print(f"  Codigo {k}: {v} veces")

# Revisar si hay codigos en blanco en filas que tienen nombre
sin_codigo = [idx+2 for idx, r in enumerate(rows_sga[1:]) if (r[0] is None or str(r[0]).strip() == '') and (r[1] is not None and str(r[1]).strip() != '')]
print(f"Filas con nombre pero SIN codigo: {len(sin_codigo)} -> Filas: {sin_codigo}")
for f in sin_codigo:
    print(f"  Fila {f}: Nombre={rows_sga[f-1][1]}")
